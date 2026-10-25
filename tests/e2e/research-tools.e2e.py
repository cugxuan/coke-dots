"""Functional image check. Run as the desktop's non-root user, without network."""
import csv
import importlib.metadata
import io
import json
import locale
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tarfile
import tempfile
import zipfile

import bs4
import docx
import pymupdf as fitz
import lxml.etree
import openpyxl
import pandas
import pypdf
import requests


def run(*args, cwd=None):
    return subprocess.run(args, cwd=cwd, check=True, text=True, capture_output=True).stdout


assert os.getuid() == 1000, "Exercise the actual Agent UID, not root"
assert sys.prefix == "/opt/coke-dots/python", sys.prefix
assert locale.setlocale(locale.LC_ALL, "") == "en_US.UTF-8"
tools = ["pdftotext", "pdfinfo", "pdftoppm", "rg", "jq", "file", "unzip", "zip", "xxd"]
assert all(shutil.which(tool) for tool in tools)
assert os.environ.get("DOTS_RESEARCH_TOOLS") == "1"
assert run("npm", "config", "get", "cache").strip() == os.environ["NPM_CONFIG_CACHE"]
for key in ["NPM_CONFIG_CACHE", "PIP_CACHE_DIR"]:
    cache = Path(os.environ[key])
    assert cache.is_relative_to("/workspace/.cache"), cache
    cache.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryFile(dir=cache) as probe:
        probe.write(b"non-root cache is writable")

with tempfile.TemporaryDirectory(prefix="dots-research-tools-") as temporary:
    root = Path(temporary)
    pdf = root / "research.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "Coke Dots research toolkit")
    document.new_page().insert_text((72, 72), "比亚迪工具验证 2026", fontname="china-s")
    document.save(pdf)
    document.close()
    extracted = root / "research.txt"
    run("pdftotext", "-layout", str(pdf), str(extracted))
    text = extracted.read_text()
    assert "Coke Dots research toolkit" in text
    assert "比亚迪工具验证" in "".join(text.split())
    assert "Pages:" in run("pdfinfo", str(pdf))
    page = run("pdftotext", "-f", "2", "-l", "2", str(pdf), "-")
    assert "Coke Dots" not in page and "比亚迪" in "".join(page.split())
    assert len(pypdf.PdfReader(pdf).pages) == 2
    assert "Coke Dots" in pypdf.PdfReader(pdf).pages[0].extract_text()
    with fitz.open(pdf) as reader:
        assert "比亚迪" in "".join(reader[1].get_text().split())
    run("pdftoppm", "-f", "1", "-singlefile", "-scale-to", "320", "-png", str(pdf), str(root / "preview"))
    assert (root / "preview.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
    assert "Coke Dots" in run("rg", "-n", "Coke Dots", str(extracted))
    assert "PDF" in run("file", str(pdf))
    assert run("xxd", "-p", "-l", "4", str(pdf)).strip() == "25504446"
    metadata = root / "metadata.json"
    metadata.write_text(json.dumps({"company": "比亚迪"}, ensure_ascii=False))
    assert run("jq", "-r", ".company", str(metadata)).strip() == "比亚迪"

    frame = pandas.DataFrame([{"公司": "比亚迪", "状态": "工具就绪", "数量": 30}])
    frame.to_excel(root / "research.xlsx", index=False, engine="openpyxl")
    assert openpyxl.load_workbook(root / "research.xlsx").active["A2"].value == "比亚迪"
    assert pandas.read_excel(root / "research.xlsx").iloc[0]["数量"] == 30
    frame.to_csv(root / "research.csv", index=False, encoding="utf-8-sig")
    with (root / "research.csv").open(encoding="utf-8-sig", newline="") as handle:
        assert list(csv.DictReader(handle))[0]["状态"] == "工具就绪"
    word = docx.Document()
    word.add_paragraph("比亚迪资料收集工具验证")
    word.save(root / "research.docx")
    assert docx.Document(root / "research.docx").paragraphs[0].text == "比亚迪资料收集工具验证"
    html = '<html><body><table><tr><td>比亚迪</td><td>工具就绪</td></tr></table></body></html>'
    assert bs4.BeautifulSoup(html, "lxml").find("td").text == "比亚迪"
    assert lxml.etree.fromstring(b"<root><value>30</value></root>").findtext("value") == "30"
    assert requests.Request("GET", "https://example.test/", params={"company": "比亚迪"}).prepare().url.endswith("company=%E6%AF%94%E4%BA%9A%E8%BF%AA")
    run("zip", "-q", str(root / "research.zip"), "research.csv", cwd=root)
    unpacked = root / "unpacked"
    run("unzip", "-q", str(root / "research.zip"), "-d", str(unpacked))
    assert (unpacked / "research.csv").read_bytes() == (root / "research.csv").read_bytes()

    # Reproduce dependency installation with the default caches and no network.
    # These fixtures are deliberately local; core research needs no installation.
    archive = root / "dots-research-fixture.tgz"
    with tarfile.open(archive, "w:gz") as package:
        for name, value in {
            "package/package.json": '{"name":"dots-research-fixture","version":"1.0.0","main":"index.cjs"}',
            "package/index.cjs": 'module.exports = "cache is writable";',
        }.items():
            payload = value.encode()
            info = tarfile.TarInfo(name)
            info.size = len(payload)
            package.addfile(info, io.BytesIO(payload))
    node_project = root / "node-project"
    node_project.mkdir()
    (node_project / "package.json").write_text('{"name":"toolkit-check","version":"1.0.0"}')
    run("npm", "install", "--offline", "--ignore-scripts", "--no-audit", "--no-fund", "--package-lock=false", str(archive), cwd=node_project)
    assert run("node", "-p", "require('dots-research-fixture')", cwd=node_project).strip() == "cache is writable"
    wheel = root / "dots_research_fixture-1.0.0-py3-none-any.whl"
    with zipfile.ZipFile(wheel, "w") as package:
        package.writestr("dots_research_fixture.py", 'value = "venv is writable"\n')
        package.writestr("dots_research_fixture-1.0.0.dist-info/METADATA", "Metadata-Version: 2.1\nName: dots-research-fixture\nVersion: 1.0.0\n")
        package.writestr("dots_research_fixture-1.0.0.dist-info/WHEEL", "Wheel-Version: 1.0\nGenerator: fixture\nRoot-Is-Purelib: true\nTag: py3-none-any\n")
        package.writestr("dots_research_fixture-1.0.0.dist-info/RECORD", "")
    run(sys.executable, "-m", "pip", "install", "--no-index", "--no-deps", "--disable-pip-version-check", str(wheel))
    assert run(sys.executable, "-c", "import dots_research_fixture; print(dots_research_fixture.value)").strip() == "venv is writable"

versions = {name: importlib.metadata.version(name) for name in ["PyMuPDF", "pypdf", "pandas", "openpyxl", "python-docx", "beautifulsoup4", "lxml", "requests"]}
print(json.dumps({"ok": True, "uid": os.getuid(), "python": sys.version.split()[0], "tools": tools, "libraries": versions, "checks": ["Chinese PDF text and page ranges", "PDF metadata and PNG", "CSV/XLSX/DOCX round trips", "HTML/XML parsing", "rg/jq/file/zip/unzip/xxd", "offline npm cache and pip venv writes"]}, ensure_ascii=False))
